"""Load the purchase scenario from JSON or an Excel workbook."""

import json
import os
from pathlib import Path


PROJECT_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_DATA_FILE = PROJECT_ROOT / "test_data" / "purchase.json"


def load_purchase_data(file_path: str | Path | None = None) -> dict:
    """Return the first purchase scenario in the selected data file."""
    selected_path = Path(file_path or os.getenv("TEST_DATA_FILE", DEFAULT_DATA_FILE))
    if not selected_path.is_absolute():
        selected_path = PROJECT_ROOT / selected_path

    suffix = selected_path.suffix.lower()
    if suffix == ".json":
        with selected_path.open(encoding="utf-8") as data_file:
            scenario = json.load(data_file)
    elif suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(selected_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value).strip() for value in next(rows)]
            values = next(rows)
            scenario = dict(zip(headers, values))
        finally:
            workbook.close()
    else:
        raise ValueError("Test data must be a .json, .xlsx, or .xlsm file")

    required_fields = {"base_url", "product_name", "quantity"}
    missing_fields = required_fields.difference(scenario)
    if missing_fields:
        raise ValueError(f"Missing test data fields: {', '.join(sorted(missing_fields))}")

    scenario["quantity"] = int(scenario["quantity"])
    if scenario["quantity"] < 1:
        raise ValueError("quantity must be greater than zero")
    return scenario